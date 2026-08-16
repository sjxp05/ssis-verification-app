from openpyxl import load_workbook

SRC = r"C:\Users\bearf\OneDrive\바탕 화면\조견표_2026.xlsx"
DST = r"C:\Users\bearf\Downloads\저장테스트.xlsx"

# 1. 원본을 열 수 있나
wb = load_workbook(SRC)
print("원본 열기 OK, 시트:", wb.sheetnames)

# 2. 아무것도 안 바꾸고 그대로 저장만
wb.save(DST)
print("저장 OK:", DST)

# 3. 저장본을 다시 열 수 있나
wb2 = load_workbook(DST)
print("재열기 OK, 시트:", wb2.sheetnames)