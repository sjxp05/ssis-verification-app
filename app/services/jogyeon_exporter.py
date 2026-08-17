#수정된 조견표 저장용도
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

#여러시트에 함께 기재되어있어서 함께 갱신해야하는 것
SYNC_KEYS = ("기본단가", "A값")

def export_modified_jogyeon(
        src_path:str|Path,
        dst_path:str|Path,
        cell_map:dict[str,tuple[str,int,int]],
        changed:dict[str,object],
)->list[tuple[str,str]]:
    #sync_keys다른 시트도 반영되도록
    changed=dict(changed)
    for key in SYNC_KEYS:
        if key in changed:
            for map_key in cell_map:
                if map_key.startswith(f"{key}@"):
                    changed.setdefault(map_key, changed[key])
            
    wb=load_workbook(src_path)
    skipped:list[tuple[str,str]]=[]

    for key,value in changed.items():
        loc=cell_map.get(key)
        if loc is None:
            skipped.append((key,"조견표 내 위치가 기록되지 않음"))
            continue
        sheet_name,r,c=loc
        if sheet_name not in wb.sheetnames:
            skipped.append((key,f"시트를 찾을 수 없음: {sheet_name}"))
            continue
        cell=wb[sheet_name].cell(row=r+1,column=c+1)
        if isinstance(cell.value,str) and cell.value.startswith("="):
            skipped.append((key,"수식으로 계산되는 셀이기에 직접 수정 불가"))
            continue
        #실제 셀값 교체
        cell.value=value

    wb.save(dst_path)
    return skipped