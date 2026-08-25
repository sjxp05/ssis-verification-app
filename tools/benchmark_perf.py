# -*- coding: utf-8 -*-
"""SSIS 검증 앱 성능 벤치마크 하네스.

기능 4종을 구간별로 계측한다:
  1) 조견표 라벨 대조  match_workbooks()
  2) 조견표 값 추출    ValueExtractor.extract_jogyeon_values()
  3) 고시법령 대조     gosi_verifier.read_gosi() (+ 원문 패널 렌더)
  4) 결제단가표 로딩   table_validator.read_prev_table()

실측 규모를 흉내 낸 합성 파일을 생성해 사용한다. 실제 파일로 재려면:
  python tools/benchmark_perf.py --baseline 작년.xlsx --target 올해.xlsx --hwpx 고시.hwpx --prev 작년단가표.xlsx

합성 데이터와 결과 로그(results.jsonl)는 tools/.bench/ 아래에 생긴다 (커밋 금지).
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "app"))

DATA = Path(__file__).resolve().parent / ".bench"
DATA.mkdir(exist_ok=True)

# ──────────────────────────────────────────────── 계측 도구
TIMES: dict[str, float] = defaultdict(float)
COUNTS: dict[str, int] = defaultdict(int)


def _reset():
    TIMES.clear()
    COUNTS.clear()


def wrap(obj, attr, label):
    orig = getattr(obj, attr)
    if getattr(orig, "_bench_wrapped", False):
        return

    def timed(*a, **k):
        t0 = time.perf_counter()
        try:
            return orig(*a, **k)
        finally:
            TIMES[label] += time.perf_counter() - t0
            COUNTS[label] += 1

    timed._bench_wrapped = True
    setattr(obj, attr, timed)


def report(feature, total, note=""):
    print(f"\n=== {feature}  총 {total:.3f}s {note}")
    for label, t in sorted(TIMES.items(), key=lambda kv: -kv[1]):
        print(f"    {label:<42} {t:8.3f}s  x{COUNTS[label]}")
    return {"feature": feature, "total": round(total, 3), "note": note,
            "sections": {k: {"t": round(v, 3), "n": COUNTS[k]} for k, v in TIMES.items()}}


# ──────────────────────────────────────────────── 합성 데이터 생성
def gen_jogyeon(path: Path, year: int, variant: bool, bloat: bool):
    """추출기 호환 레이아웃 + 대량 필러 표. variant=True 면 라벨 일부 변형(올해 파일)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bump = 1.03 if variant else 1.0

    # 인정조사 ------------------------------------------------------
    ws = wb.active
    ws.title = "인정조사"
    ws["B1"] = "기본단가"; ws["C1"] = int(17270 * bump)
    ws["B2"] = "A값"; ws["C2"] = int(3089062 * bump); ws["D2"] = int(154400 * bump)
    ws["B3"] = "기본 부담률"
    for i, v in enumerate([0.06, 0.09, 0.12, 0.15]):
        ws.cell(3, 3 + i, v)
    for i, v in enumerate([0.02, 0.03, 0.04, 0.05]):
        ws.cell(4, 3 + i, v)
    ws["B6"] = "활동지원등급"; ws["C6"] = "지원시간"; ws["D6"] = "상한시간"
    ws["E6"] = "기본형 한도"; ws["F6"] = "확장형 한도"
    for i in range(4):
        r = 8 + i
        ws.cell(r, 2, f"{i + 1}등급")
        ws.cell(r, 3, round(118.2 - 23 * i, 1))
        ws.cell(r, 4, 130)
        ws.cell(r, 5, int((2041000 - 400000 * i) * bump))
        ws.cell(r, 6, int((1841000 - 400000 * i) * bump))

    _fill_bulk(ws, start_row=14, variant=variant, prefix="소득적용")

    # 산정특례 ------------------------------------------------------
    ws = wb.create_sheet("산정특례")
    ws["B1"] = int(154400 * bump); ws["C1"] = "본인부담금 상한액"
    ws["B2"] = "기준중위소득"
    for i, v in enumerate([0.06, 0.09, 0.12, 0.15]):
        ws.cell(3, 2 + i, v)
    from app.config.anchors import ADD_ITEMS
    for i, item in enumerate(ADD_ITEMS):
        ws.cell(5, 2 + i, item)
        ws.cell(6, 2 + i, int((3000000 - 150000 * i) * bump))
    _fill_bulk(ws, start_row=9, variant=variant, prefix="특례구간")

    # 종합조사 ------------------------------------------------------
    ws = wb.create_sheet("종합조사")
    ws["A1"] = "주간활동 기본형"
    for i in range(15):
        ws.cell(1, 2 + i, f"{i + 1}등급")
        ws.cell(2, 2 + i, int((5940000 - 405000 * i) * bump))
    ws["A4"] = "주간활동 확장형"
    for i in range(15):
        ws.cell(4, 2 + i, f"{i + 1}등급")
        ws.cell(5, 2 + i, int((5840000 - 405000 * i) * bump))
    _fill_bulk(ws, start_row=8, variant=variant, prefix="점수구간")

    if bloat:  # 서식만 있는 먼 셀 → used range 부풀림 재현
        for s in wb.worksheets:
            s.cell(3000, 30).font = Font(bold=True)

    wb.save(path)


BULK_BLOCKS = 3
BULK_ROWS = 40
BULK_COLS = 13


def _fill_bulk(ws, start_row: int, variant: bool, prefix: str):
    """실제 조견표의 큰 본문 표를 흉내: 텍스트 라벨 + 숫자 행렬."""
    r = start_row
    for block in range(BULK_BLOCKS):
        ws.cell(r, 2, f"{prefix} 제{block + 1}편 산출 기준표")
        r += 1
        for i in range(BULK_ROWS):
            n = block * BULK_ROWS + i
            label = f"{prefix} {n + 1}차 적용대상"
            if variant:
                if n % 12 == 5:
                    label = f"{prefix} {n + 1}차 적용 대상(개정)"   # 개칭 → 하이브리드 유도
                elif n % 12 == 9:
                    label = f"{n + 1}차 {prefix} 대상"              # 어순 변경
            ws.cell(r, 2, label)
            for c in range(3, BULK_COLS):
                ws.cell(r, c, (n + 1) * 1000 + c)
            r += 1
        if variant and block == 0:
            ws.cell(r, 2, f"{prefix} 신설항목")                     # 신규 라벨
            for c in range(3, BULK_COLS):
                ws.cell(r, c, 999)
            r += 1
        r += 2  # 빈 줄로 표 분리


def gen_hwpx(path: Path):
    """gosi_verifier 앵커를 모두 만족하는 합성 고시 + 대량 필러."""
    from xml.sax.saxutils import escape

    NS = 'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"'

    def p(text):
        return f"<hp:p><hp:run><hp:t>{escape(text)}</hp:t></hp:run></hp:p>"

    def tbl(rows):
        out = ["<hp:tbl>"]
        for row in rows:
            out.append("<hp:tr>")
            for cell in row:
                out.append(f"<hp:tc><hp:subList><hp:p><hp:run><hp:t>{escape(str(cell))}</hp:t></hp:run></hp:p></hp:subList></hp:tc>")
            out.append("</hp:tr>")
        out.append("</hp:tbl>")
        return "".join(out)

    won = lambda n: f"{n:,}원"
    parts = [f'<?xml version="1.0" encoding="UTF-8"?><hp:sec {NS}>']

    # 필러 앞부분 (제1~2장)
    parts.append(p("제1장 총칙"))
    for i in range(150):
        parts.append(p(f"제{i + 1}조의 세부 시행 기준에 관한 설명 문단이다. " * 4))
    for t in range(10):
        parts.append(p(f"{t + 10}. 참고자료"))
        parts.append(tbl([[f"항목{r}-{c}" for c in range(4)] for r in range(20)]))

    # 제3장 — 서비스 단가
    parts.append(p("제3장 급여비용 및 산정기준"))
    parts.append(p("1. 활동보조"))
    parts.append(tbl([
        ["구분", "시간당 금액", "가산수당"],
        ["일반(평일 주간)", won(16150), won(1010)],
        ["심야(22시~06시)", won(24220), won(1510)],
        ["공휴일 및 근로자의 날", won(24220), won(1510)],
    ]))
    parts.append(p("2. 방문목욕"))
    parts.append(tbl([
        ["구분", "금액(1회당)"],
        ["이동목욕용 차량내에서 목욕을 제공한 경우", won(84240)],
        ["가정내에서 목욕을 제공한 경우", won(75960)],
    ]))
    parts.append(p("3. 방문간호"))
    parts.append(tbl([
        ["구분", "금액"],
        ["30분미만", won(40760)],
        ["30분이상 60분미만", won(51110)],
        ["60분이상", won(61490)],
    ]))
    parts.append(p("4. 방문간호지시서"))
    parts.append(tbl([
        ["발급기관", "발급 경우", "금액"],
        ["「의료법」에 따른 의료기관", "대상자가 의료기관을 방문하여 발급, 의사가 가정을 방문하여 발급", f"{won(20790)} {won(72740)}"],
        ["「지역보건법」에 따른 보건소", "대상자가 보건소를 방문하여 발급, 의사가 가정을 방문하여 발급", f"{won(15790)} {won(57740)}"],
    ]))
    parts.append(p("5. 월 한도액"))
    parts.append(tbl(
        [["활동지원급여 구간", "종합점수", "금액"]]
        + [[f"{i + 1}구간", f"{465 - 30 * i}점 이상", won(8472000 - 405000 * i)] for i in range(15)]
    ))
    parts.append(tbl(
        [["구간", "주간활동 기본형", "주간활동 확장형"]]
        + [[f"{i + 1}구간", won(5940000 - 405000 * i), won(5840000 - 405000 * i)] for i in range(15)]
    ))

    # 필러 중간
    for i in range(150):
        parts.append(p(f"경과 규정 및 적용 예시에 대한 부가 설명 문단 {i}이다. " * 4))

    # 부칙
    parts.append(p("부 칙"))
    parts.append(tbl([
        ["활동지원등급", "1등급", "2등급", "3등급", "4등급"],
        ["기본급여", won(2041000), won(1641000), won(1241000), won(841000)],
    ]))
    parts.append(tbl([
        ["추가급여 대상", "금액"],
        ["독거 가구로 종합점수 400점이상", won(2732000)],
        ["독거 가구로 종합점수 380점미만", won(200000)],
        ["가구구성원이 있는 취약가구로 400점이상", won(1092000)],
        ["수급자 본인이 출산한 경우", won(800000)],
        ["자립을 준비하는 경우", won(200000)],
        ["학교에 다니는 경우", won(100000)],
        ["직장에 다니는 경우", won(400000)],
        ["보호자가 일시적으로 부재한 경우", won(200000)],
        ["나머지 가구구성원의 직장생활 등", won(730000)],
    ]))
    for i in range(100):
        parts.append(p(f"부칙 시행일 및 경과조치 설명 문단 {i}이다. " * 3))
    parts.append("</hp:sec>")

    import zipfile
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("Contents/section0.xml", "".join(parts))


def gen_price_table(path: Path, rows: int = 1500):
    """결제단가표 흉내: 스타일 입힌 수천 행짜리 표."""
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, Side

    wb = Workbook()
    ws = wb.active
    thin = Border(*(Side(style="thin"),) * 4)
    headers = ["코드", "등급구분", "지원량", "정부지원금", "본인부담금", "합계"]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(1, c).font = Font(bold=True)
    tiers = "가나다라마바"
    for i in range(rows):
        tier = tiers[i % 6]
        ws.append([f"S{i:05d}", f"{i % 15 + 1}구간({tier}형)", 100 + i % 60,
                   3000000 - (i % 100) * 7000, (i % 100) * 3000, 3000000 - (i % 100) * 4000])
        for c in range(1, len(headers) + 1):
            ws.cell(i + 2, c).border = thin
            ws.cell(i + 2, c).number_format = "#,##0"
    wb.save(path)


# ──────────────────────────────────────────────── 벤치마크 본체
def bench_match(baseline, target, runs=2):
    from jogyeon_matcher import engine
    from jogyeon_matcher.ingest import loader, segmenter
    from jogyeon_matcher import encoder, similarity
    from jogyeon_matcher.validation import value_checks

    wrap(loader, "load", "loader.load (엑셀 적재+정제)")
    wrap(loader, "_sanitize", "loader._sanitize")
    wrap(segmenter, "find_tables", "segmenter.find_tables")
    wrap(segmenter, "compare_sheets", "segmenter.compare_sheets")
    wrap(segmenter, "check_anchor_uniqueness", "segmenter.check_anchor_uniqueness")
    wrap(value_checks, "compare_regions", "value_checks.compare_regions")
    wrap(engine, "_collect_labels", "engine._collect_labels")
    wrap(
        engine,
        "_match_deterministic",
        "label_rules.match_deterministic (절대식 매칭)",
    )
    wrap(engine, "_find_missing", "engine._find_missing")
    wrap(encoder.DenseEncoder, "_load", "DenseEncoder._load (모델 로드)")
    wrap(encoder.DenseEncoder, "encode", "DenseEncoder.encode (임베딩)")
    wrap(similarity.BM25, "scores", "BM25.scores")

    out = []
    for i in range(runs):
        _reset()
        t0 = time.perf_counter()
        rep = engine.match_workbooks(baseline, target)
        total = time.perf_counter() - t0
        s = rep.summary
        out.append(
            report(
                "1. 조견표 라벨 대조",
                total,
                f"(run{i + 1}) 라벨 {s.total_labels} "
                f"auto {s.auto_passed} review {s.needs_review} "
                f"unmatched {s.unmatched}",
            )
        )
    return out


def bench_extract(sheet_path, runs=2):
    import pandas as pd
    from app.models.dto import UploadedFile
    from app.services import jogyeon_value_extractor as jve

    wrap(jve.ValueExtractor, "_load", "ValueExtractor._load (read_excel+정규화사본)")
    wrap(jve.ValueExtractor, "_find_all", "ValueExtractor._find_all (키워드 전셀검색)")
    wrap(pd, "read_excel", "pd.read_excel")

    out = []
    for i in range(runs):
        _reset()
        t0 = time.perf_counter()
        data = jve.ValueExtractor().extract_jogyeon_values(UploadedFile.from_path(sheet_path))
        total = time.perf_counter() - t0
        out.append(report("2. 조견표 값 추출", total, f"(run{i + 1}) 키 {sum(len(d) for d in data)}개"))
    return out


def bench_gosi(hwpx_path, runs=2):
    from app.services import gosi_verifier as gv

    wrap(gv, "parse_document", "parse_document (hwpx XML 파싱)")
    wrap(gv, "find_table", "find_table (앵커 탐색)")
    wrap(gv, "validate_prices", "validate_prices")
    wrap(gv, "compare_with_reference", "compare_with_reference")

    out = []
    result = None
    for i in range(runs):
        _reset()
        t0 = time.perf_counter()
        result = gv.read_gosi(str(hwpx_path), reference=None)
        total = time.perf_counter() - t0
        n_issue = len(result["검증"])
        out.append(report("3a. 고시 파싱+대조(read_gosi)", total,
                          f"(run{i + 1}) 블록 {len(result['블록'])} 이슈 {n_issue}"))

    # UI 렌더 (offscreen)
    try:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PyQt6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication([])
        from app.ui.widgets.gosi_panel import GosiDocumentViewer

        viewer = GosiDocumentViewer()
        _reset()
        t0 = time.perf_counter()
        try:
            viewer.set_blocks(result["블록"], chapters=("제3장",))  # 단일 렌더 경로
            style = "단일렌더"
        except TypeError:
            viewer.set_blocks(result["블록"])          # 구버전: 이중 렌더
            viewer.set_visible_chapters(("제3장",))
            style = "이중렌더"
        t2 = time.perf_counter()
        TIMES[f"렌더 ({style})"] = t2 - t0; COUNTS[f"렌더 ({style})"] = 1
        out.append(report("3b. 고시 원문 패널 렌더", t2 - t0, f"(offscreen, {style})"))
    except Exception as error:
        print(f"    [3b 렌더 측정 실패: {error}]")
    return out


def bench_prev_table(path, runs=2):
    from app.services import table_validator as tv

    out = []
    for i in range(runs):
        _reset()
        t0 = time.perf_counter()
        df = tv.read_prev_table(str(path))
        total = time.perf_counter() - t0
        out.append(report("4. 결제단가표(작년) 로딩", total, f"(run{i + 1}) {df.shape[0]}행 x {df.shape[1]}열"))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--baseline"); ap.add_argument("--target")
    ap.add_argument("--hwpx"); ap.add_argument("--prev")
    ap.add_argument("--bloat", action="store_true", help="used-range 부풀림 변형 사용")
    ap.add_argument("--big", action="store_true", help="대용량 변형(시트당 수천 행)")
    ap.add_argument("--only", choices=["match", "extract", "gosi", "prev"])
    ap.add_argument("--tag", default="baseline")
    args = ap.parse_args()

    global BULK_BLOCKS, BULK_ROWS, BULK_COLS
    suffix = "_bloat" if args.bloat else ""
    if args.big:
        BULK_BLOCKS, BULK_ROWS, BULK_COLS = 8, 250, 28   # 시트당 약 2000행 x 28열
        suffix += "_big"
    baseline = Path(args.baseline) if args.baseline else DATA / f"jogyeon_2026{suffix}.xlsx"
    target = Path(args.target) if args.target else DATA / f"jogyeon_2027{suffix}.xlsx"
    hwpx = Path(args.hwpx) if args.hwpx else DATA / "gosi_synthetic.hwpx"
    prev = Path(args.prev) if args.prev else DATA / f"price_table_prev{suffix}.xlsx"

    if not args.baseline and not baseline.exists():
        print("합성 조견표 생성 중...")
        gen_jogyeon(baseline, 2026, variant=False, bloat=args.bloat)
        gen_jogyeon(target, 2027, variant=True, bloat=args.bloat)
    if not args.hwpx and not hwpx.exists():
        print("합성 고시(hwpx) 생성 중...")
        gen_hwpx(hwpx)
    if not args.prev and not prev.exists():
        print("합성 결제단가표 생성 중...")
        gen_price_table(prev, rows=6000 if args.big else 1500)

    results = []
    if args.only in (None, "match"):
        results += bench_match(baseline, target)
    if args.only in (None, "extract"):
        results += bench_extract(baseline)
    if args.only in (None, "gosi"):
        results += bench_gosi(hwpx)
    if args.only in (None, "prev"):
        results += bench_prev_table(prev)

    log = DATA / "results.jsonl"
    with log.open("a", encoding="utf-8") as f:
        for r in results:
            r["tag"] = args.tag
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    print(f"\n결과 {len(results)}건 기록: {log}")


if __name__ == "__main__":
    main()
