# 라벨 매칭 테스트 케이스 생성기 -> synthetic_label_tests.xlsx
#
#   python tests/fixtures/generate_label_tests.py
#   python tests/fixtures/generate_label_tests.py --source 실제_조견표.xlsx
#
# --source 를 주면 실제 조견표에서 앵커가 부분문자열로 몇 번 나오는지 스캔해
# '앵커중복스캔' 시트를 함께 만든다. 없어도 테스트케이스 생성에는 지장이 없다.
#
# 케이스를 늘리려면 아래 변형 사전에 항목만 추가하면 된다.
# 기대결과의 의미는 README 시트에 적혀 있고, 판정은 tests/test_label_matching.py 가 한다.

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
from typing import NamedTuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parent / "synthetic_label_tests.xlsx"

# ── 기준 라벨 인벤토리 — 조견표와 단가표에서 실제로 쓰는 문구 ──────────────
ANCHORS = [
    "기본단가", "A값", "기본 부담률", "상한액", "활동지원등급", "기준중위소득",
    "주간활동 기본형", "주간활동 확장형",
]
IJ_GRADES = [f"{i}등급" for i in range(1, 5)]
JH_LABELS = [f"{i}구간" for i in range(1, 16)]
ADD_ITEMS = [
    "최중증1인가구", "1등급1인가구", "2등급이하1인가구", "최중증취약가구",
    "1등급취약가구", "2등급이하취약가구", "출산", "자립준비", "학교생활",
    "직장생활", "보호자일시부재", "나머지가구구성원의직장생활등",
]
TABLE_HEADERS = [
    "안", "사업유형ID", "사업년도", "차수", "등급구분", "등급명", "바우처구분",
    "지원량", "정부지원금", "본인부담금", "재판정여부", "재판정기간",
    "소득기준ID", "소득기준년도", "소득구분", "정렬순서", "물품코드",
]
IJ_INCOME = [
    "기초수급자", "차상위", "전국가구평균소득50%이하",
    "전국가구평균소득50%초과~100%이하", "전국가구평균소득100%초과~150%이하",
    "전국가구평균소득150%초과~200%이하",
]
JH_INCOME = [
    "기초수급자", "차상위", "기준중위소득70%이하", "기준중위소득70%초과~120%이하",
    "기준중위소득120%초과~180%이하", "기준중위소득180%초과",
]
EXTRACTED_KEYS = [
    "기본단가", "종합조사/산정특례 본인부담금 상한액", "종합조사/산정특례 본인부담률",
    "인정조사 월한도액 (기본형)", "추가급여 월한도액",
]
SHEET_BANDS = ["50% 이하", "100% 이하", "150% 이하", "150% 초과"]

ALL_BASE = sorted(set(
    ANCHORS + IJ_GRADES + JH_LABELS + ADD_ITEMS + TABLE_HEADERS
    + IJ_INCOME + JH_INCOME + EXTRACTED_KEYS + SHEET_BANDS
))

# ── 변형 사전 — 케이스를 늘리려면 여기만 손대면 된다 ──────────────────────
# A2: 문자 겹침이 적어 임베딩이 견인해야 하는 도메인 동의어
SYNONYM = {
    "정부지원금": ["국고지원금", "정부 지원액"],
    "본인부담금": ["자기부담금", "자부담금"],
    "지원량": ["급여량", "지원 시간"],
    "등급명": ["등급 명칭"],
    "소득구분": ["소득유형"],
    "재판정여부": ["재판정 대상 여부"],
    "물품코드": ["품목코드"],
    "정렬순서": ["표시순서"],
    "기초수급자": ["기초생활수급자"],
    "차상위": ["차상위계층"],
    "학교생활": ["학교 재학"],
    "직장생활": ["직장 재직"],
    "출산": ["출산가구"],
    "나머지가구구성원의직장생활등": ["가족의직장생활"],
    "기본단가": ["시간당 단가"],
    "상한액": ["한도 금액"],
}

# A3: 줄이거나 늘려 쓴 표현
EXPAND = {
    "주간활동 기본형": ["주간활동(기본)", "주간활동 서비스 기본형"],
    "주간활동 확장형": ["주간활동(확장)", "주간활동 서비스 확장형"],
    "보호자일시부재": ["보호자 일시 부재", "보호자의 일시적 부재"],
    "활동지원등급": ["활동지원 등급", "장애인활동지원 등급"],
    "본인부담금": ["본인부담금액", "수급자 본인부담금"],
    "기본 부담률": ["기본부담률(%)", "기본 본인부담률"],
    "A값": ["A값(전국평균소득)", "2026 A값"],
    "사업년도": ["사업 연도"],
    "소득기준년도": ["소득기준 연도"],
}

# A4: 어순이나 조사가 바뀐 표현
REORDER = {
    "종합조사/산정특례 본인부담금 상한액": ["본인부담금 상한액(종합조사·산정특례)"],
    "인정조사 월한도액 (기본형)": ["기본형 월한도액(인정조사)"],
    "추가급여 월한도액": ["월한도액(추가급여)"],
    "주간활동 기본형": ["기본형 주간활동"],
    "기준중위소득70%이하": ["중위소득 기준 70% 이하"],
}

# A5: 자소 단위 오타
TYPO = {
    "기본단가": "기본단까", "본인부담금": "본인부덤금", "차상위": "차상휘",
    "물품코드": "물풍코드", "학교생활": "학교샹활", "지원량": "지원랑",
}

# A1 표기 변형을 적용할 라벨
SURFACE_TARGETS = [
    "기본단가", "본인부담금", "주간활동 기본형", "기본 부담률",
    "지원량", "활동지원등급", "상한액",
]

# B2: 조견표와 무관한 문자열. 전 후보가 컷오프 미만이어야 한다.
NEGATIVES = [
    "김치찌개", "서울특별시 강남구", "OperatingSystem",
    "감가상각누계액", "휴대폰 요금제", "π=3.14159",
]

# C2: 형제 밴드가 문자·의미 모두 근접해 숫자만이 판별자인 집합
BAND_FAMILIES = {
    "인정조사 소득밴드": IJ_INCOME[2:],
    "종합조사 소득밴드": JH_INCOME[2:],
    "구간(1~15)": JH_LABELS,
    "등급(1~4)": IJ_GRADES,
    "시트 헤더 밴드": SHEET_BANDS,
}

# C3: 임베딩 유사도가 비정상적으로 높은 반의어 쌍
ANTONYMS = [
    ("150% 이하", "150% 초과"),
    ("기준중위소득70%이하", "기준중위소득180%초과"),
    ("주간활동 기본형", "주간활동 확장형"),
    ("전국가구평균소득50%이하", "전국가구평균소득150%초과~200%이하"),
]

# C1: 과잉 정규화 시 서로 다른 기준 라벨이 합쳐지는지 확인하는 프로브
COLLISION_PROBES = [
    ("1등급", "2등급", "숫자 제거형 정규화 금지"),
    ("50% 이하", "100% 이하", "숫자·% 제거형 정규화 금지"),
    ("기본 부담률", "추가 부담률", "수식어 제거형 정규화 금지"),
    ("주간활동 기본형", "주간활동 확장형", "접미 유형어 제거형 정규화 금지"),
]

# C4: 눈에 보이지 않는 문자. no_match 나 오매칭으로 새면 실패다.
INVISIBLES = [
    ("A값", "A​값", "ZWSP(U+200B)"),
    ("A값", "Ａ값", "전각 A(U+FF21) — NFKC로 해소"),
    ("기본단가", "기본⁠단가", "WORD JOINER(U+2060)"),
    ("본인부담금", "본인부담금﻿", "BOM(U+FEFF)"),
    ("차상위", "차상위 ", "NBSP(U+00A0)"),
]

# C5: 짧은 앵커가 긴 라벨의 부분문자열이라 첫 매칭 위치에 의존하면 오염되는 경우
SUBSTRING_TRAPS = [
    ("1등급", ["1등급1인가구", "1등급취약가구"], "앵커가 다른 라벨의 접두 부분문자열"),
    ("상한액", ["종합조사/산정특례 본인부담금 상한액"], "짧은 앵커 ⊂ 긴 라벨"),
    ("활동지원등급", ["등급구분", "등급명"], "'등급' 공유 — 산정특례 시트 3회 출현"),
    ("주간활동 확장형", ["주간활동 기본형"], "종합조사 시트 2회 출현 — 첫 매칭 위치 의존 금지"),
    ("기본단가", ["시간당 단가"], "4개 시트 동시 출현 — 시트 스코프 필수"),
]

# C6: 라벨이 같은데 의미만 바뀐 경우. 텍스트로는 원리상 못 잡는 사각지대.
BLINDSPOTS = [
    ("주간활동 기본형", "라벨은 동일하나 열 위치·의미가 바뀐 경우 — 값 수준 체크 필요"),
    ("지원시간", "단위 변경(시간→분)도 라벨이 같으면 통과 — 값 범위 체크로만 방어 가능"),
]

JUDGING_RULES = [
    ("auto_pass", "입력==기준(원문 그대로). 1순위 & 점수≈1.0 아니면 파이프라인 버그"),
    ("auto_pass_after_norm", "normalize(입력)==normalize(기준) 이어야 함. 실패 시 정규화 결함"),
    ("review_top3", "기준라벨이 후보 상위 3위 안에 있어야 함 (Recall@3)"),
    ("top1_strict", "기준라벨이 1순위 & '지정 교란 후보'들보다 점수가 높아야 함"),
    ("pairwise_distinct", "normalize(기준) != normalize(교란 후보) — 정규화 충돌 금지"),
    ("never_silent_fail", "auto_pass 또는 기준이 top3 면 통과. no_match·오매칭은 실패"),
    ("no_match", "모든 후보 점수가 컷오프 미만이어야 함"),
    ("auto_pass_BLINDSPOT", "통과가 정상이나 라벨 검증의 구조적 사각지대. 값 수준 체크로 별도 방어"),
]

COLUMNS = ["test_id", "카테고리", "세부유형", "기준라벨(정답)", "입력라벨(변형)",
           "지정 교란 후보", "기대결과", "검증 대상 컴포넌트", "비고"]


class Case(NamedTuple):
    category: str
    subtype: str
    base: str
    variant: str
    expected: str
    component: str
    note: str = ""
    distractors: str = ""


def surface_variants(label: str) -> list[str]:
    """A1: 단위 주석·공백·전각·장식 괄호. 정규화가 흡수해야 한다."""
    spaced = (
        label.replace(" ", "") if " " in label
        else label[: len(label) // 2] + " " + label[len(label) // 2 :]
    )
    bracketed = (
        label.replace("(", "（").replace(")", "）") if "(" in label else f"[{label}]"
    )
    unit = f"{label}(단위: 원)" if label.endswith("액") else f"{label}(원)"
    return [unit, spaced, f" {label} ", bracketed]


def build_cases() -> list[Case]:
    cases: list[Case] = []

    # B1 항등 — 자기 자신이 1순위가 아니면 배선이 잘못된 것이다
    cases += [
        Case("B.메타모픽", "B1.항등", b, b, "auto_pass",
             "파이프라인 배선(프리픽스·정규화 버그)", "자기 자신 → 1순위·점수≈1 필수")
        for b in ALL_BASE
    ]

    cases += [
        Case("B.메타모픽", "B2.음성대조군", "(없음)", word, "no_match",
             "컷오프·점수 스케일", "모든 후보 점수 < 컷오프 필수")
        for word in NEGATIVES
    ]

    cases += [
        Case("A.일반변형", "A1.표기(공백·괄호·단위·전각)", base, variant,
             "auto_pass_after_norm", "정규화 + BM25", "정규화 후 완전일치 → 검토 큐 진입 금지")
        for base in SURFACE_TARGETS
        for variant in surface_variants(base)
    ]

    cases += [
        Case("A.일반변형", "A2.동의어", base, variant, "review_top3", "임베딩",
             "BM25 토큰 겹침 낮음 → 임베딩이 상위 3위 내로 올려야 함")
        for base, variants in SYNONYM.items()
        for variant in variants
    ]

    cases += [
        Case("A.일반변형", "A3.축약확장", base, variant, "review_top3", "BM25+임베딩 합산")
        for base, variants in EXPAND.items()
        for variant in variants
    ]

    cases += [
        Case("A.일반변형", "A4.어순·조사", base, variant, "review_top3",
             "토큰화(n-gram)+임베딩")
        for base, variants in REORDER.items()
        for variant in variants
    ]

    cases += [
        Case("A.일반변형", "A5.오타·자소", base, variant, "review_top3", "문자 n-gram")
        for base, variant in TYPO.items()
    ]

    # C2 숫자 판별자 — 형제 밴드를 교란 후보로 붙여 정답이 1순위인지 본다
    for family in BAND_FAMILIES.values():
        for i, base in enumerate(family):
            siblings = [x for j, x in enumerate(family) if j != i][:4]
            variant = (
                base.replace("~", " ~ ").replace("%", "% ") if "%" in base
                else base[:-1] + " " + base[-1]
            )
            cases.append(Case(
                "C.우회·적대", "C2.숫자판별자", base, variant, "top1_strict",
                "숫자 토큰 가중(BM25·임베딩 공통 약점)",
                "형제 밴드가 문자·의미 모두 초근접 — 숫자만이 판별자",
                " | ".join(siblings),
            ))

    for base, opposite in ANTONYMS:
        variant = base.replace("이하", " 이하").replace("초과", " 초과").replace("형", " 형")
        cases.append(Case(
            "C.우회·적대", "C3.반의어·부정", base, variant, "top1_strict",
            "임베딩 반의어 맹점",
            "이하↔초과·기본↔확장은 임베딩 유사도가 비정상적으로 높음", opposite,
        ))

    for base, other, why in COLLISION_PROBES:
        cases.append(Case(
            "C.우회·적대", "C1.정규화충돌", base, base, "pairwise_distinct",
            "정규화 함수 자체", f"norm({base!r}) != norm({other!r}) 이어야 함 — {why}", other,
        ))

    for base, variant, why in INVISIBLES:
        cases.append(Case(
            "C.우회·적대", "C4.비가시문자", base, variant, "never_silent_fail",
            "유니코드 정규화 범위",
            f"{why} — auto_pass 또는 review 허용, no_match·오매칭 금지",
        ))

    for base, traps, why in SUBSTRING_TRAPS:
        cases.append(Case(
            "C.우회·적대", "C5.부분문자열함정", base, base, "top1_strict",
            "완전일치 우선순위·앵커 유일성", why, " | ".join(traps),
        ))

    cases += [
        Case("C.우회·적대", "C6.동일문자열_의미이동", base, base, "auto_pass_BLINDSPOT",
             "라벨 수준 검증 불가", why)
        for base, why in BLINDSPOTS
    ]

    return cases


def scan_anchors(source: Path) -> list[tuple[str, int, str, str]]:
    """실제 조견표에서 앵커가 부분문자열로 몇 번 나오는지 센다 (C5 근거)."""
    targets = sorted(set(ANCHORS + ["월 한도액", "본인부담금", "지원시간", "1구간"] + IJ_GRADES))
    hits = {a: Counter() for a in targets}

    workbook = load_workbook(source, read_only=True)
    for name in workbook.sheetnames:
        for row in workbook[name].iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str):
                    for anchor in targets:
                        if anchor in cell:
                            hits[anchor][name] += 1

    report = []
    for anchor in targets:
        total = sum(hits[anchor].values())
        per_sheet = ", ".join(f"{k}:{v}" for k, v in hits[anchor].items())
        if total > 1 or len(hits[anchor]) > 1:
            risk = "위험: 유일성 미보장 — 시트 스코프·완전일치 우선 필수"
        elif total == 1:
            risk = "안전(현재 파일 기준)"
        else:
            risk = "미출현 — 앵커 자체 재검토"
        report.append((anchor, total, per_sheet, risk))
    return report


BOLD = Font(name="Arial", bold=True)
PLAIN = Font(name="Arial")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
RISK_FILL = PatternFill("solid", fgColor="FCE4EC")


def _write_header(sheet, columns, widths):
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = BOLD
        cell.fill = HEAD_FILL
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, i).column_letter].width = width
    sheet.freeze_panes = "A2"


def write_xlsx(cases: list[Case], out: Path, scan: list | None = None) -> None:
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "테스트케이스"
    _write_header(sheet, COLUMNS, [8, 12, 24, 30, 32, 34, 20, 26, 44])
    for i, case in enumerate(cases, 1):
        sheet.append([
            f"T{i:04d}", case.category, case.subtype, case.base, case.variant,
            case.distractors, case.expected, case.component, case.note,
        ])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = PLAIN
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if str(row[1].value).startswith("C."):
            for cell in row:
                cell.fill = RISK_FILL

    if scan:
        anchor_sheet = workbook.create_sheet("앵커중복스캔")
        _write_header(anchor_sheet, ["앵커", "부분문자열 총 출현", "시트별 출현", "위험도 판정"],
                      [28, 16, 60, 44])
        for anchor, total, per_sheet, risk in scan:
            anchor_sheet.append([anchor, total, per_sheet, risk])
            for cell in anchor_sheet[anchor_sheet.max_row]:
                cell.font = PLAIN
                if total != 1:
                    cell.fill = RISK_FILL

    readme = workbook.create_sheet("README")
    readme.append(["기대결과 판정 규칙 (tests/test_label_matching.py 가 자동 판정)"])
    readme["A1"].font = BOLD
    for name, rule in JUDGING_RULES:
        readme.append([name, rule])
    readme.append([])
    readme.append(["케이스를 추가하려면 generate_label_tests.py 의 변형 사전에 항목을 넣고 다시 실행"])
    for row in readme.iter_rows(min_row=2):
        for cell in row:
            cell.font = PLAIN
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 100

    workbook.save(out)


def main() -> None:
    parser = argparse.ArgumentParser(description="라벨 매칭 테스트 케이스 생성")
    parser.add_argument("--source", type=Path, help="실제 조견표 xlsx (앵커 중복 스캔용)")
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args()

    cases = build_cases()
    scan = scan_anchors(args.source) if args.source else None
    write_xlsx(cases, args.out, scan)

    print(f"{len(cases)}개 케이스 -> {args.out}")
    for subtype, count in sorted(Counter(c.subtype for c in cases).items()):
        print(f"  {subtype:<28} {count:>3}")
    if scan is None:
        print("\n앵커중복스캔 시트 없음 (--source 로 실제 조견표를 주면 생성)")


if __name__ == "__main__":
    main()
